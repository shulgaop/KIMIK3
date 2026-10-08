#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
============================================================
 ОТЧЁТ О ПОХОДЕ: GPS из фото и видео -> карта + таблица
============================================================

Что делает скрипт:
  1. Сканирует папку (и все подпапки) с фотографиями и видео.
  2. Извлекает GPS-координаты и дату/время съёмки:
     - фото — из EXIF;
     - видео с телефона — из метаданных контейнера (атом ©xyz);
     - видео с DJI Osmo Action 4/5/6 — из встроенной телеметрии
       (через библиотеку pyosmogps); извлекается весь GPS-трек
       и рисуется на карте пунктирной линией;
     - файлы БЕЗ координат — привязываются к GPS-треку: если в папке
       лежит GPX-файл (трек с часов, навигатора, Strava, Maps.me),
       фото/видео ставится в точку трека, ближайшую по времени съёмки.
       Сам трек рисуется на карте отдельным слоем.
       Для видео DJI время съёмки берётся из имени файла
       (DJI_20260906095623_... = 06.09.2026 09:56:23), поэтому привязка
       к треку работает даже без телеметрии.
       Если за день нет записанного трека, но есть плановый маршрут
       (GPX без меток времени) — позиция оценивается интерполяцией
       по нему пропорционально времени (помечается «по плану (оценка)»).
  3. Определяет название местности (офлайн-база, интернет не нужен).
  4. Строит маршрут по дням и создаёт три файла:
     - карта_похода.html     — интерактивная карта: линии маршрута
       по дням, треки DJI, маркеры со всплывающими превью фото;
     - отчёт_похода.xlsx     — таблица: файл (кликабельно), дата,
       координаты (кликабельная ссылка на Google Maps), местность,
       день маршрута;
     - файлы_без_геометок.txt — список того, что привязать не удалось;
     - tiles_offline/        — локальные тайлы карты Esri (OFFLINE_MAP=True):
       с ними карта работает полностью без интернета. Интернет нужен
       один раз — на момент запуска скрипта.

Установка зависимостей (один раз):
    pip install pillow folium openpyxl reverse-geocoder
    pip install pyosmogps        # для видео с DJI Osmo Action
    pip install pillow-heif      # только если есть HEIC-фото с iPhone

Запуск:
    python hike_report.py                # обработать текущую папку
    python hike_report.py "D:\\Поход"    # или указать папку явно

Если в папке есть видео DJI, при запуске будет предложен выбор:
искать GPS в телеметрии видео, привязать к GPX-треку по времени
или использовать оба способа.

ВАЖНО про DJI Osmo Action 6:
  У камеры НЕТ встроенного GPS. Координаты в видео появляются, только
  если во время съёмки камера была подключена к GPS Bluetooth-пульту
  DJI (или приложению, передающему GPS с телефона). Без этого видео
  попадут в список "без геометок" — это ограничение камеры.
  Фото, присланные через WhatsApp/Telegram, обычно лишены EXIF:
  используйте оригиналы файлов.
"""

import os
import re
import io
import sys
import mmap
import base64
import subprocess
import importlib.util
import html as html_mod
import xml.etree.ElementTree as ET
from datetime import datetime, timedelta

PHOTO_EXT = {'.jpg', '.jpeg', '.png', '.tif', '.tiff', '.webp', '.heic', '.heif'}
VIDEO_EXT = {'.mp4', '.mov', '.m4v', '.avi', '.mkv', '.mts', '.3gp'}

EPOCH_1904 = datetime(1904, 1, 1)  # точка отсчёта времени в MP4/MOV

# Часовой пояс места похода: в видео DJI и в GPX время хранится в UTC,
# а в фото — локальное. Для Фанских гор (Таджикистан) это UTC+5.
# Для похода в другом поясе поменяйте обе цифры (Москва = 3).
DJI_TIMEZONE_OFFSET = 5

# --- Офлайн-карта ---
# HIKE_OFFLINE_MAP=0 в окружении отключает скачивание тайлов (так делает
# веб-мастер — иначе шаг «Фото» ждёт тайлы несколько минут)
OFFLINE_MAP = os.environ.get("HIKE_OFFLINE_MAP", "1") != "0"
                          # True — скачать локальные тайлы и добавить
                          # офлайн-слой (интернет нужен один раз, при запуске)
OFFLINE_MIN_ZOOM = 10     # минимальный масштаб тайлов
OFFLINE_MAX_ZOOM = 14     # максимальный масштаб; автоматически уменьшится,
                          # если тайлов получится слишком много
MAX_OFFLINE_TILES = 4000  # защитный лимит на число скачиваемых тайлов

# Источник офлайн-тайлов — топографическая карта Esri.
# Можно заменить на спутник:
# 'https://server.arcgisonline.com/ArcGIS/rest/services/'
# 'World_Imagery/MapServer/tile/{z}/{y}/{x}'
OFFLINE_TILE_URL = ('https://server.arcgisonline.com/ArcGIS/rest/services/'
                    'World_Topo_Map/MapServer/tile/{z}/{y}/{x}')

# --- Привязка файлов без координат к GPS-треку (GPX) ---
# Просто положите GPX-файл в папку с фото — скрипт подхватит его сам.
GPX_TIMEZONE_OFFSET = 5   # часовой пояс: в GPX время хранится в UTC,
                          # а в фото — локальное (Фаны/Таджикистан = 5)
TRACK_MAX_GAP_MIN = 30    # привязывать фото к треку, только если точка
                          # трека по времени не дальше этих минут от съёмки
PLAN_INTERPOLATION = True  # дни без записанного трека: оценивать позицию
                          # по плановому маршруту (GPX без меток времени),
                          # распределяя фото пропорционально времени между
                          # опорными точками записанных треков

# Поддержка HEIC (iPhone), если установлен pillow-heif
try:
    import pillow_heif
    pillow_heif.register_heif_opener()
except ImportError:
    pass

HAVE_RG = importlib.util.find_spec('reverse_geocoder') is not None
HAVE_PYOSMO = importlib.util.find_spec('pyosmogps') is not None

if not HAVE_RG:
    print('[!] reverse_geocoder не установлен — названия мест определяться не будут.')
    print('    Установите: pip install reverse-geocoder')
if not HAVE_PYOSMO:
    print('[!] pyosmogps не установлен — GPS из видео DJI Osmo Action извлекаться не будет.')
    print('    Установите: pip install pyosmogps')


# ------------------------------------------------------------------
#  Извлечение данных из фото (EXIF)
# ------------------------------------------------------------------
def _dms_to_deg(dms, ref):
    """Градусы/минуты/секунды из EXIF -> десятичные градусы."""
    d, m, s = (float(x) for x in dms)
    val = d + m / 60.0 + s / 3600.0
    if ref in ('S', 'W'):
        val = -val
    return val


def photo_meta(path):
    """Возвращает (lat, lon, dt) для фотографии."""
    from PIL import Image
    lat = lon = dt = None
    with Image.open(path) as img:
        exif = img.getexif()

        # Дата/время съёмки: DateTimeOriginal -> DateTimeDigitized -> DateTime
        for ifd_tag, key in ((0x8769, 36867), (0x8769, 36868), (None, 306)):
            try:
                src = exif.get_ifd(ifd_tag) if ifd_tag else exif
            except Exception:
                continue
            v = src.get(key)
            if v:
                try:
                    dt = datetime.strptime(str(v).strip(), '%Y:%m:%d %H:%M:%S')
                    break
                except ValueError:
                    pass

        # GPS-блок
        try:
            gps = exif.get_ifd(0x8825)
        except Exception:
            gps = {}
        if gps and 2 in gps and 4 in gps:
            try:
                lat = _dms_to_deg(gps[2], str(gps.get(1, 'N')))
                lon = _dms_to_deg(gps[4], str(gps.get(3, 'E')))
            except Exception:
                lat = lon = None
    return lat, lon, dt


# ------------------------------------------------------------------
#  Извлечение данных из видео с телефона (атомы MP4/MOV)
# ------------------------------------------------------------------
ISO6709_RE = re.compile(rb'([+-]\d{2}\.\d{3,})([+-]\d{3}\.\d{3,})')


def video_meta(path):
    """Возвращает (lat, lon, dt) для видео. Ищет атом ©xyz и время mvhd."""
    lat = lon = dt = None
    with open(path, 'rb') as f:
        mm = mmap.mmap(f.fileno(), 0, access=mmap.ACCESS_READ)
    try:
        # Координаты: атом ©xyz (ISO 6709, напр. "+55.7558+037.6173/")
        idx = mm.find(b'\xa9xyz')
        if idx != -1:
            m = ISO6709_RE.search(mm[idx:idx + 200])
            if m:
                lat, lon = float(m.group(1)), float(m.group(2))

        # Дата создания: атом mvhd
        idx = mm.find(b'mvhd')
        if idx != -1 and idx + 20 < len(mm):
            version = mm[idx + 4]
            if version == 1:
                secs = int.from_bytes(mm[idx + 8:idx + 16], 'big')
            else:
                secs = int.from_bytes(mm[idx + 8:idx + 12], 'big')
            if 0 < secs < 4_000_000_000:
                # mvhd хранит время в UTC — переводим в локальное,
                # чтобы совпадало с EXIF фото и сдвинутым GPX-треком
                dt = (EPOCH_1904 + timedelta(seconds=secs)
                      + timedelta(hours=GPX_TIMEZONE_OFFSET))
    finally:
        mm.close()
    return lat, lon, dt


DJI_NAME_RE = re.compile(r'DJI_(\d{14})', re.IGNORECASE)


def dt_from_filename(path):
    """Дата/время съёмки из имени файла DJI: DJI_20260906095623_0004_D.MP4
    -> 2026-09-06 09:56:23 (время локальное, как в EXIF фото)."""
    m = DJI_NAME_RE.search(os.path.basename(path))
    if not m:
        return None
    try:
        return datetime.strptime(m.group(1), '%Y%m%d%H%M%S')
    except ValueError:
        return None


# ------------------------------------------------------------------
#  Извлечение GPS-трека из видео DJI Osmo Action (pyosmogps)
# ------------------------------------------------------------------
def parse_gpx(path):
    """Читает GPX-файл ->
    (точки_с_временем [(dt, lat, lon)],
     точки_без_времени [(lat, lon)],      # плановый маршрут
     путевые_точки [(имя, lat, lon)]).    # wpt: лагеря, перевалы и т.п."""
    timed, untimed, wpts = [], [], []
    tree = ET.parse(path)
    for p in tree.iter():
        if p.tag.endswith('trkpt'):
            try:
                lat, lon = float(p.get('lat')), float(p.get('lon'))
            except (TypeError, ValueError):
                continue
            dt = None
            for child in p:
                if child.tag.endswith('time') and child.text:
                    try:
                        dt = datetime.fromisoformat(
                            child.text.strip().replace('Z', '+00:00')
                        ).replace(tzinfo=None)
                    except ValueError:
                        pass
            if dt:
                timed.append((dt, lat, lon))
            else:
                untimed.append((lat, lon))
        elif p.tag.endswith('wpt'):
            try:
                lat, lon = float(p.get('lat')), float(p.get('lon'))
            except (TypeError, ValueError):
                continue
            name = ''
            for child in p:
                if child.tag.endswith('name') and child.text:
                    name = child.text.strip()
            wpts.append((name, lat, lon))
    return timed, untimed, wpts


def nearest_track_point(track, dt, max_gap_min):
    """Ближайшая по времени точка трека к моменту dt (бинарный поиск).
    track — отсортированный список (dt, lat, lon).
    Возвращает (lat, lon, разница_в_минутах) или None."""
    import bisect
    if not track or dt is None:
        return None
    times = [p[0] for p in track]
    i = bisect.bisect_left(times, dt)
    best = None
    for j in (i - 1, i):
        if 0 <= j < len(track):
            gap = abs((track[j][0] - dt).total_seconds())
            if best is None or gap < best[0]:
                best = (gap, track[j])
    if best and best[0] <= max_gap_min * 60:
        return best[1][1], best[1][2], best[0] / 60.0
    return None


def dji_track(path):
    """Извлекает GPS-трек из видео DJI через pyosmogps.
    Возвращает список (dt, lat, lon) или None, если трека нет."""
    gpx_path = path + '.tmp.gpx'
    try:
        cmd = [sys.executable, '-m', 'pyosmogps',
               '-t', str(DJI_TIMEZONE_OFFSET),
               '-f', '1', '-r', 'lpf',          # 1 точка/сек, сглаживание
               'extract', path, gpx_path]
        r = subprocess.run(cmd, capture_output=True, text=True, timeout=3600)
        if r.returncode != 0 or not os.path.exists(gpx_path):
            return None
        pts, _, _ = parse_gpx(gpx_path)
        return pts or None
    except Exception:
        return None
    finally:
        if os.path.exists(gpx_path):
            try:
                os.remove(gpx_path)
            except OSError:
                pass


# ------------------------------------------------------------------
#  Локальные тайлы для офлайн-карты
# ------------------------------------------------------------------
def _deg2num(lat, lon, z):
    """Координаты -> номер тайла (схема slippy map)."""
    import math
    n = 2 ** z
    x = int((lon + 180.0) / 360.0 * n)
    lat_r = math.radians(max(-85.0511, min(85.0511, lat)))
    y = int((1.0 - math.asinh(math.tan(lat_r)) / math.pi) / 2.0 * n)
    return x, y


def download_offline_tiles(records, folder):
    """Скачивает тайлы (источник — OFFLINE_TILE_URL), покрывающие весь
    маршрут с запасом. Возвращает (успех, число_тайлов, макс_zoom).
    Уже скачанные тайлы пропускаются — повторный запуск продолжит
    с места обрыва."""
    import urllib.request

    lats = [r['lat'] for r in records]
    lons = [r['lon'] for r in records]
    pad_lat = (max(lats) - min(lats)) * 0.15 + 0.01
    pad_lon = (max(lons) - min(lons)) * 0.15 + 0.01
    lat_min, lat_max = min(lats) - pad_lat, max(lats) + pad_lat
    lon_min, lon_max = min(lons) - pad_lon, max(lons) + pad_lon

    def tile_ranges(z):
        x1, y1 = _deg2num(lat_max, lon_min, z)
        x2, y2 = _deg2num(lat_min, lon_max, z)
        return (range(min(x1, x2), max(x1, x2) + 1),
                range(min(y1, y2), max(y1, y2) + 1))

    # Подбираем максимальный zoom, чтобы уложиться в лимит тайлов
    max_z = OFFLINE_MAX_ZOOM
    while max_z > OFFLINE_MIN_ZOOM:
        total = sum(len(tile_ranges(z)[0]) * len(tile_ranges(z)[1])
                    for z in range(OFFLINE_MIN_ZOOM, max_z + 1))
        if total <= MAX_OFFLINE_TILES:
            break
        max_z -= 1

    jobs = [(z, x, y)
            for z in range(OFFLINE_MIN_ZOOM, max_z + 1)
            for x in tile_ranges(z)[0]
            for y in tile_ranges(z)[1]]

    print(f'Загрузка офлайн-тайлов: {len(jobs)} шт. '
          f'(zoom {OFFLINE_MIN_ZOOM}–{max_z})...')
    os.makedirs(folder, exist_ok=True)
    done = 0
    try:
        for z, x, y in jobs:
            dst = os.path.join(folder, str(z), str(x), f'{y}.png')
            if os.path.exists(dst):
                done += 1
                continue
            os.makedirs(os.path.dirname(dst), exist_ok=True)
            url = OFFLINE_TILE_URL.format(z=z, x=x, y=y)
            req = urllib.request.Request(
                url, headers={'User-Agent': 'hike-report/1.0 (offline tiles)'})
            with urllib.request.urlopen(req, timeout=20) as resp:
                data = resp.read()
            # Отбрасываем всё, что не похоже на PNG/JPEG (заглушки
            # вида "API key required", JSON-ошибки сервера и т.п.)
            if not (data.startswith(b'\x89PNG') or data.startswith(b'\xff\xd8')):
                continue
            with open(dst, 'wb') as f:
                f.write(data)
            done += 1
            if done % 200 == 0:
                print(f'  тайлов загружено: {done}/{len(jobs)}', flush=True)
    except Exception as e:
        print(f'  загрузка тайлов прервана: {e}')
    if done < len(jobs):
        print(f'  внимание: скачано {done} из {len(jobs)} тайлов — '
              f'офлайн-слой будет с пропусками. Перезапустите скрипт, '
              f'чтобы докачать (скачанное сохраняется).')
    return done > 0, done, max_z


# ------------------------------------------------------------------
#  Вспомогательные функции
# ------------------------------------------------------------------
def coords_ok(lat, lon):
    return (lat is not None and lon is not None
            and -90 <= lat <= 90 and -180 <= lon <= 180
            and not (lat == 0 and lon == 0))


def thumbnail_b64(path, max_px=320):
    """Маленькое превью фото в base64 для всплывающего окна на карте."""
    from PIL import Image
    try:
        with Image.open(path) as img:
            img.thumbnail((max_px, max_px))
            buf = io.BytesIO()
            img.convert('RGB').save(buf, 'JPEG', quality=70)
        return base64.b64encode(buf.getvalue()).decode()
    except Exception:
        return None


def scan_folder(root):
    files = []
    for dirpath, _, names in os.walk(root):
        for n in sorted(names):
            ext = os.path.splitext(n)[1].lower()
            if ext in PHOTO_EXT or ext in VIDEO_EXT:
                files.append(os.path.join(dirpath, n))
    return files


def nearest_plan_idx(plan, lat, lon):
    """Индекс ближайшей точки планового маршрута к (lat, lon)."""
    best_i, best_d = 0, float('inf')
    for i, (la, lo) in enumerate(plan):
        dd = (la - lat) ** 2 + (lo - lon) ** 2
        if dd < best_d:
            best_d, best_i = dd, i
    return best_i


def build_plan_anchors(plan, track, pending_dts):
    """Опорные точки [(время, индекс_на_плане)] для интерполяции:
    начало и конец каждого записанного сегмента трека привязываются
    к ближайшей точке планового маршрута. Края (до первого / после
    последнего трека) достраиваются по времени непривязанных файлов."""
    anchors = []
    if track:
        # разбиваем записанный трек на сегменты (разрыв > 3 ч)
        cuts = [0]
        for i in range(1, len(track)):
            if (track[i][0] - track[i - 1][0]).total_seconds() > 3 * 3600:
                cuts.append(i)
        cuts.append(len(track))
        for a, b in zip(cuts, cuts[1:]):
            for dt, lat, lon in (track[a], track[b - 1]):
                anchors.append((dt, float(nearest_plan_idx(plan, lat, lon))))
    before = [d for d in pending_dts if not anchors or d < anchors[0][0]]
    after = [d for d in pending_dts if not anchors or d > anchors[-1][0]]
    if before:
        anchors.insert(0, (min(before) - timedelta(hours=1), 0.0))
    if after:
        anchors.append((max(after) + timedelta(hours=1), float(len(plan) - 1)))
    return sorted(anchors)


def interp_plan(plan, anchors, dt):
    """Позиция на плановом маршруте для момента dt: линейная
    интерполяция между опорными точками. None, если вне интервалов."""
    for (t0, i0), (t1, i1) in zip(anchors, anchors[1:]):
        if t0 <= dt <= t1 and t1 > t0:
            frac = (dt - t0).total_seconds() / (t1 - t0).total_seconds()
            pos = i0 + frac * (i1 - i0)
            i = int(pos)
            j = min(i + 1, len(plan) - 1)
            f2 = pos - i
            lat = plan[i][0] + (plan[j][0] - plan[i][0]) * f2
            lon = plan[i][1] + (plan[j][1] - plan[i][1]) * f2
            return lat, lon
    return None


def ask_dji_mode():
    """Интерактивный выбор способа обработки видео DJI.
    1 — только телеметрия, 2 — только по треку, 3 — комбо (по умолчанию)."""
    print('\nКак обрабатывать видео DJI?')
    print('  1 — искать GPS в телеметрии видео (точно, но медленно;')
    print('      нужен pyosmogps и съёмка с GPS-пультом)')
    print('  2 — привязать к GPX-треку по времени съёмки (быстро;')
    print('      нужен .gpx-файл в папке)')
    print('  3 — сначала телеметрия, если её нет — по треку')
    try:
        ans = input('Ваш выбор [1/2/3, Enter = 3]: ').strip()
    except (EOFError, KeyboardInterrupt):
        ans = ''
    mode = {'1': 1, '2': 2}.get(ans, 3)
    print('Режим DJI:', {1: 'телеметрия', 2: 'по треку',
                         3: 'телеметрия + трек'}[mode])
    return mode


# ------------------------------------------------------------------
#  Основная логика
# ------------------------------------------------------------------
def main():
    # python hike_report.py [папка] [--gpx файл1 файл2 ...]
    # --gpx: треки из проекта (записанный с часов, плановый маршрут) — читаются
    # в дополнение к GPX из папки с фото; записанный приоритетнее планового.
    argv = sys.argv[1:]
    extra_gpx = []
    if "--gpx" in argv:
        i = argv.index("--gpx")
        extra_gpx = [a for a in argv[i + 1:] if not a.startswith("--")]
        argv = argv[:i]
    root = os.path.abspath(argv[0] if argv else os.getcwd())
    print(f'Папка: {root}')

    files = scan_folder(root)
    if not files:
        print('Фото и видео не найдены. Проверьте путь к папке.')
        return
    print(f'Найдено файлов: {len(files)}.')

    # --- Загрузка GPS-треков: все GPX-файлы из папки и подпапок + --gpx ---
    gpx_points = []      # записанный трек (с метками времени)
    plan_points = []     # плановый маршрут (без времени)
    waypoints = []       # путевые точки (wpt)
    gpx_files = [os.path.join(dp, n) for dp, _, names in os.walk(root)
                 for n in sorted(names) if n.lower().endswith('.gpx')]
    gpx_files += [g for g in extra_gpx if os.path.isfile(g)]
    for gpx_path in gpx_files:
        n = os.path.basename(gpx_path)
        timed, untimed, wpts = parse_gpx(gpx_path)
        gpx_points.extend(timed)
        plan_points.extend(untimed)
        waypoints.extend(wpts)
        desc = f'  трек: {n} — {len(timed)} точек с временем'
        if untimed:
            desc += f' + {len(untimed)} плановых'
        if wpts:
            desc += f' + {len(wpts)} путевых'
        print(desc)
    # Время в GPX — UTC, переводим в локальное (как в EXIF фото)
    gpx_points = [(dt + timedelta(hours=GPX_TIMEZONE_OFFSET) if dt else None,
                   la, lo) for dt, la, lo in gpx_points]
    gpx_points = sorted((p for p in gpx_points if p[0] is not None),
                        key=lambda p: p[0])
    if gpx_points:
        print(f'GPS-трек загружен: {len(gpx_points)} точек, '
              f'{gpx_points[0][0]:%d.%m %H:%M} — {gpx_points[-1][0]:%d.%m %H:%M}.')
        print('Файлы без координат будут привязаны к треку по времени съёмки.')
    else:
        print('GPX-трек не найден (положите .gpx в папку, '
              'чтобы привязать файлы без координат).')

    # --- Выбор режима обработки видео DJI ---
    has_dji = any(os.path.splitext(f)[1].lower() in VIDEO_EXT
                  and DJI_NAME_RE.search(os.path.basename(f))
                  for f in files)
    dji_mode = ask_dji_mode() if has_dji else 3
    if has_dji:
        if dji_mode in (1, 3) and not HAVE_PYOSMO:
            print('[!] Телеметрия недоступна: pyosmogps не установлен '
                  '(pip install pyosmogps).')
        if dji_mode in (2, 3) and not gpx_points:
            print('[!] Выбрана привязка к треку, но GPX не найден — '
                  'видео DJI без телеметрии останутся без координат.')

    print('Обработка...')
    records, no_gps, pending = [], [], []
    n_dji = n_track = 0
    for i, path in enumerate(files, 1):
        ext = os.path.splitext(path)[1].lower()
        kind = 'фото' if ext in PHOTO_EXT else 'видео'
        try:
            lat, lon, dt = photo_meta(path) if kind == 'фото' else video_meta(path)
        except Exception:
            lat = lon = dt = None
        # У DJI дата/время съёмки зашиты в имя файла (DJI_20260906095623_...)
        # — это надёжнее метаданных mvhd
        dt = dt_from_filename(path) or dt

        # Видео DJI: если в контейнере координат нет — пробуем телеметрию
        # (режимы 1 и 3)
        is_dji = bool(DJI_NAME_RE.search(os.path.basename(path)))
        track = None
        source = 'EXIF' if kind == 'фото' else 'метаданные'
        if (kind == 'видео' and is_dji and not coords_ok(lat, lon)
                and dji_mode in (1, 3) and HAVE_PYOSMO):
            print(f'  [{i}/{len(files)}] {os.path.basename(path)} — '
                  f'поиск телеметрии DJI...', flush=True)
            pts = dji_track(path)
            if pts:
                track = [(la, lo) for _, la, lo in pts]
                lat, lon = pts[0][1], pts[0][2]   # позиция видео = старт трека
                if dt is None:
                    dt = next((p[0] for p in pts if p[0]), None)
                source = 'DJI-трек'
                n_dji += 1

        # Нет координат ни в файле, ни в телеметрии — привязка к GPX-треку
        # (для видео DJI — только в режимах 2 и 3)
        skip_track = is_dji and dji_mode == 1
        if (not coords_ok(lat, lon) and gpx_points and dt is not None
                and not skip_track):
            m = nearest_track_point(gpx_points, dt, TRACK_MAX_GAP_MIN)
            if m:
                lat, lon, gap_min = m
                source = f'по треку (±{gap_min:.0f} мин)'
                n_track += 1

        rel = os.path.relpath(path, root)
        rec = {
            'file': rel, 'abs': path, 'kind': kind, 'source': source,
            'lat': lat, 'lon': lon, 'dt': dt, 'place': '', 'track': track,
        }
        if coords_ok(lat, lon):
            if dt is None:  # запасной вариант — время изменения файла
                rec['dt'] = datetime.fromtimestamp(os.path.getmtime(path))
            records.append(rec)
        elif dt is not None:
            pending.append(rec)   # время есть — попробуем плановый маршрут
        else:
            no_gps.append(rel)
        status = (f'OK ({lat:.5f}, {lon:.5f}) [{source}]'
                  if coords_ok(lat, lon)
                  else ('нет GPS — проверю по плановому маршруту'
                        if dt is not None else 'нет GPS'))
        print(f'  [{i}/{len(files)}] {rel} — {status}')

    # --- Дни без записанного трека: оценка позиции по плановому маршруту ---
    n_plan = 0
    if pending and plan_points and PLAN_INTERPOLATION:
        anchors = build_plan_anchors(plan_points, gpx_points,
                                     [r['dt'] for r in pending])
        for r in pending:
            pos = interp_plan(plan_points, anchors, r['dt'])
            if pos:
                r['lat'], r['lon'] = pos
                r['source'] = 'по плану (оценка)'
                records.append(r)
                n_plan += 1
                print(f'  {r["file"]} — по плану ({pos[0]:.5f}, {pos[1]:.5f})')
            else:
                no_gps.append(r['file'])
        if n_plan:
            print(f'По плановому маршруту расставлено файлов: {n_plan} '
                  f'(оценочные позиции!)')
    else:
        no_gps.extend(r['file'] for r in pending)

    if not records:
        print('\nНи в одном файле не найдено координат.')
        print('Частые причины: фото из мессенджеров (EXIF вырезан); '
              'у DJI Action не был подключён GPS-пульт; '
              'выключена геометка в камере телефона.')
        print('Если у вас есть записанный трек (часы, навигатор, Strava, '
              'Maps.me) — экспортируйте его в GPX и положите в папку: '
              'скрипт привяжет файлы по времени съёмки.')
        return

    # --- Названия мест (офлайн) ---
    if HAVE_RG:
        import reverse_geocoder as rg
        print('Определение названий мест...')
        geo = rg.search([(r['lat'], r['lon']) for r in records])
        for r, g in zip(records, geo):
            parts = [g.get('name', ''), g.get('admin1', '')]
            r['place'] = ', '.join(p for p in parts if p) or '—'

    # --- Разбивка по дням ---
    records.sort(key=lambda r: r['dt'])
    days = sorted({r['dt'].date() for r in records})
    day_of = {d: i + 1 for i, d in enumerate(days)}
    for r in records:
        r['day'] = day_of[r['dt'].date()]

    # --- Интерактивная карта ---
    import folium
    center = [sum(r['lat'] for r in records) / len(records),
              sum(r['lon'] for r in records) / len(records)]

    # Офлайн-тайлы: скачиваем заранее, пока есть интернет
    offline_ok, off_max_z = False, OFFLINE_MAX_ZOOM
    if OFFLINE_MAP:
        offline_ok, n_tiles, off_max_z = download_offline_tiles(
            records, os.path.join(root, 'tiles_offline'))
        if offline_ok:
            print(f'Офлайн-тайлы готовы: {n_tiles} шт. в папке tiles_offline')
        else:
            print('Офлайн-тайлы не загружены (нет интернета?) — '
                  'карта будет использовать онлайн-подложки.')

    # Подложки. По умолчанию — Esri World Topo (топографическая карта,
    # работает из локальных файлов без ключа и без блокировок).
    # OSM/OpenTopoMap оставлены переключаемыми слоями, но из локального
    # файла их серверы могут блокировать (ошибка 403).
    fmap = folium.Map(location=center, zoom_start=11, tiles=None)
    if offline_ok:
        folium.TileLayer(
            tiles='tiles_offline/{z}/{x}/{y}.png',
            attr='Esri, USGS, NOAA and the GIS User Community',
            name='Офлайн-карта (без интернета)',
            max_native_zoom=off_max_z, max_zoom=18,
            show=True).add_to(fmap)
    folium.TileLayer(
        tiles='https://server.arcgisonline.com/ArcGIS/rest/services/'
              'World_Topo_Map/MapServer/tile/{z}/{y}/{x}',
        attr='Esri, USGS, NOAA and the GIS User Community',
        name='Карта (Esri Topo)', show=not offline_ok).add_to(fmap)
    folium.TileLayer(
        tiles='https://server.arcgisonline.com/ArcGIS/rest/services/'
              'World_Imagery/MapServer/tile/{z}/{y}/{x}',
        attr='Esri, Maxar, Earthstar Geographics',
        name='Спутник (Esri)', show=False).add_to(fmap)
    folium.TileLayer(
        tiles='https://{s}.tile.opentopomap.org/{z}/{x}/{y}.png',
        attr='© OpenStreetMap contributors, SRTM | © OpenTopoMap (CC-BY-SA)',
        name='OpenTopoMap', show=False).add_to(fmap)
    folium.TileLayer(
        tiles='https://tile.openstreetmap.org/{z}/{x}/{y}.png',
        attr='© OpenStreetMap contributors',
        name='OpenStreetMap', show=False).add_to(fmap)

    # Загруженный GPX-трек — отдельным слоем под маршрутами дней
    if gpx_points:
        line = [(la, lo) for _, la, lo in gpx_points]
        step = max(1, len(line) // 3000)
        g = folium.FeatureGroup(name='GPS-трек (GPX)')
        folium.PolyLine(line[::step], color='black', weight=3,
                        opacity=0.7, dash_array='8').add_to(g)
        g.add_to(fmap)

    # Плановый маршрут (GPX без меток времени) и путевые точки
    if plan_points or waypoints:
        g = folium.FeatureGroup(name='Плановый маршрут (GPX)')
        if plan_points:
            step = max(1, len(plan_points) // 3000)
            folium.PolyLine(plan_points[::step], color='darkgreen', weight=3,
                            opacity=0.6, dash_array='3').add_to(g)
        for name, la, lo in waypoints:
            folium.CircleMarker(
                (la, lo), radius=5, color='darkgreen',
                fill=True, fill_opacity=0.9,
                tooltip=html_mod.escape(name) if name else None,
            ).add_to(g)
        g.add_to(fmap)
    colors = ['red', 'blue', 'green', 'purple', 'orange', 'darkred',
              'cadetblue', 'darkgreen', 'darkpurple', 'pink']

    for day in day_of.values():
        pts = [r for r in records if r['day'] == day]
        color = colors[(day - 1) % len(colors)]
        group = folium.FeatureGroup(name=f'День {day} ({pts[0]["dt"]:%d.%m.%Y})')

        # Линия маршрута дня по всем точкам
        folium.PolyLine([(r['lat'], r['lon']) for r in pts],
                        color=color, weight=4, opacity=0.8).add_to(group)

        # Треки DJI (пунктир), прореженные до ~600 точек
        for r in pts:
            if r['track']:
                step = max(1, len(r['track']) // 600)
                folium.PolyLine(r['track'][::step], color=color, weight=2,
                                opacity=0.6, dash_array='5').add_to(group)

        for r in pts:
            title = html_mod.escape(os.path.basename(r['file']))
            place = html_mod.escape(r['place'])
            body = (f'<b>{title}</b><br>{r["dt"]:%d.%m.%Y %H:%M}<br>'
                    f'{place}<br><i>{r["source"]}</i><br>')
            if r['kind'] == 'фото':
                b64 = thumbnail_b64(r['abs'])
                if b64:
                    body += f'<img src="data:image/jpeg;base64,{b64}" width="300">'
            icon = 'camera' if r['kind'] == 'фото' else 'facetime-video'
            folium.Marker(
                location=(r['lat'], r['lon']),
                popup=folium.Popup(body, max_width=320),
                icon=folium.Icon(color=color, icon=icon, prefix='fa'),
            ).add_to(group)
        group.add_to(fmap)

    folium.LayerControl(collapsed=False).add_to(fmap)
    map_path = os.path.join(root, 'карта_похода.html')
    fmap.save(map_path)

    # --- Таблица Excel ---
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = 'Маршрут'
    headers = ['№', 'Файл', 'Тип', 'День', 'Дата и время',
               'Широта', 'Долгота', 'Координаты (ссылка)', 'Местность', 'Источник']
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True)

    for i, r in enumerate(records, 1):
        url = f'https://www.google.com/maps?q={r["lat"]:.6f},{r["lon"]:.6f}'
        ws.append([i, r['file'], r['kind'], r['day'],
                   r['dt'].strftime('%d.%m.%Y %H:%M'),
                   round(r['lat'], 6), round(r['lon'], 6),
                   f'{r["lat"]:.6f}, {r["lon"]:.6f}', r['place'], r['source']])
        row = i + 1
        # кликабельное имя файла (откроется, если xlsx лежит рядом с фото)
        ws.cell(row=row, column=2).hyperlink = r['file']
        ws.cell(row=row, column=2).style = 'Hyperlink'
        # кликабельные координаты -> Google Maps
        ws.cell(row=row, column=8).hyperlink = url
        ws.cell(row=row, column=8).style = 'Hyperlink'

    widths = [5, 40, 8, 6, 18, 11, 11, 22, 30, 12]
    for idx, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(idx)].width = w
    ws.freeze_panes = 'A2'

    xlsx_path = os.path.join(root, 'отчёт_похода.xlsx')
    wb.save(xlsx_path)

    # --- Файлы без геометок ---
    txt_path = os.path.join(root, 'файлы_без_геометок.txt')
    with open(txt_path, 'w', encoding='utf-8') as f:
        f.write('В этих файлах координаты не найдены:\n\n')
        f.write('\n'.join(no_gps) if no_gps else 'Таких файлов нет.')

    # --- Итог ---
    print('\n========== ГОТОВО ==========')
    print(f'Обработано файлов:        {len(files)}')
    print(f'С координатами:           {len(records)}'
          f' (DJI-треки: {n_dji}, по GPX-треку: {n_track}, '
          f'по плану: {n_plan})')
    print(f'Без координат:            {len(no_gps)}')
    print(f'Дней в маршруте:          {len(days)}')
    print(f'Карта:                    {map_path}')
    print(f'Таблица:                  {xlsx_path}')
    print(f'Список без геометок:      {txt_path}')
    if offline_ok:
        print(f'Офлайн-тайлы:             {os.path.join(root, "tiles_offline")}')
        print('   При переносе карты на другой компьютер копируйте '
              'карта_похода.html вместе с папкой tiles_offline.')


if __name__ == '__main__':
    main()
    # Пауза, чтобы окно не закрывалось при запуске двойным кликом
    try:
        if sys.stdin.isatty():
            input('\nНажмите Enter для выхода...')
    except (EOFError, KeyboardInterrupt):
        pass
